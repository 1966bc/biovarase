# -*- coding: utf-8 -*-
# -----------------------------------------------------------------------------
# project:  biovarase
# authors:  1966bc
# mailto:   [giuseppecostanzi@gmail.com]
# modify:   winter MMXXV - hierarchical treeview (workstations → results)
# -----------------------------------------------------------------------------
"""
Daily QC Validation module.

Hierarchical TreeView:
- Parent nodes: Workstations with aggregated counts
- Child nodes: Individual results (loaded on expand)

Features:
    - Role-based validation control (Admin/Superuser can validate)
    - Workstation-level approval with audit trail (daily_approvals table)
    - Color coding: green=approved, yellow=pending, red=problems
    - Expand/collapse to drill down into results
"""

import sys
import threading
import tkinter as tk

from i18n import _
from views.parent_view import ParentView
from tkinter import ttk
from tkinter import messagebox
from datetime import datetime
from calendarium import Calendarium


# Role constants
ROLE_ADMIN = 0
ROLE_SUPERUSER = 1
ROLE_TECHNICIAN = 2
ROLE_AUTOLOGIN = 3

# Node type tags
TAG_WORKSTATION = "ws"
TAG_RESULT = "result"


class UI(ParentView):
    """
    Daily QC Validation window with hierarchical TreeView.

    Workstation nodes show:
        - Name, equipment, result counts, approval status

    Result nodes (children) show:
        - Test method, batch, result value, z-score, validation status
    """

    def __init__(self, parent):
        """Initialize the validation window."""
        super().__init__(parent, name="daily_validation")

        if self._reusing:
            return

        self.title(_("Daily QC Validation"))
        self.protocol("WM_DELETE_WINDOW", self.on_close)
        self.bind("<Escape>", self.on_close)
        self.bind("<F5>", lambda e: self._load_data())

        # Internal state
        self.selected_date = None
        self.can_validate = False
        self._refresh_job = None  # Auto-refresh timer

        # Data dictionaries
        self.dict_workstations = {}  # item_id -> workstation data
        self.dict_results = {}       # item_id -> result data
        self.loaded_ws = set()       # workstation_ids already loaded

        # Build UI
        self._build_ui()
        self.show(on_screen=True)

    def _build_ui(self):
        """Build the complete UI."""
        paddings = {"padx": 5, "pady": 5}

        # Main frame
        self.frm_main = ttk.Frame(self, style="App.TFrame", padding=8)
        self.frm_main.pack(fill=tk.BOTH, expand=True)

        # Top frame: filters
        frm_top = ttk.Frame(self.frm_main, style="App.TFrame")
        frm_top.pack(side=tk.TOP, fill=tk.X, **paddings)

        # Date selector
        ttk.Label(frm_top, text=_("Date:")).pack(side=tk.LEFT, **paddings)

        bg = self.engine.get_rgb(240, 240, 237)
        self.calendarium = Calendarium(frm_top, "", base_bg_color=bg)
        self.calendarium.pack(side=tk.LEFT, **paddings)
        self.calendarium.set_today()

        # Load button (Alt-C)
        self.btn_load = ttk.Button(
            frm_top,
            text=_("Load"),
            command=self._load_data,
            underline=0
        )
        self.btn_load.pack(side=tk.LEFT, padx=(10, 0))
        self.bind("<Alt-c>", lambda e: self._load_data())

        # Role indicator
        self.lbl_role = ttk.Label(frm_top, text="", foreground="blue")
        self.lbl_role.pack(side=tk.RIGHT, **paddings)

        # Middle frame: TreeView
        frm_tree = ttk.Frame(self.frm_main, style="App.TFrame")
        frm_tree.pack(side=tk.TOP, fill=tk.BOTH, expand=True, **paddings)

        # Scrollbars
        sb_vert = ttk.Scrollbar(frm_tree, orient=tk.VERTICAL)
        sb_horiz = ttk.Scrollbar(frm_tree, orient=tk.HORIZONTAL)

        # TreeView with hierarchical structure
        cols = ("equipment_batch", "time", "counts_result", "problems_sd", "status")
        self.tree = ttk.Treeview(
            frm_tree,
            columns=cols,
            yscrollcommand=sb_vert.set,
            xscrollcommand=sb_horiz.set,
            height=18
        )
        sb_vert.config(command=self.tree.yview)
        sb_horiz.config(command=self.tree.xview)

        # Column headers
        self.tree.heading("#0", text=_("Workstation / Test"), anchor=tk.W)
        self.tree.heading("equipment_batch", text=_("Equipment / Batch"), anchor=tk.W)
        self.tree.heading("time", text=_("Time"), anchor=tk.CENTER)
        self.tree.heading("counts_result", text=_("Counts / Result"), anchor=tk.CENTER)
        self.tree.heading("problems_sd", text=_("Problems / Z-Score"), anchor=tk.CENTER)
        self.tree.heading("status", text=_("Status"), anchor=tk.CENTER)

        # Column widths
        self.tree.column("#0", width=200, anchor=tk.W)
        self.tree.column("equipment_batch", width=140, anchor=tk.W)
        self.tree.column("time", width=50, anchor=tk.CENTER)
        self.tree.column("counts_result", width=100, anchor=tk.CENTER)
        self.tree.column("problems_sd", width=80, anchor=tk.CENTER)
        self.tree.column("status", width=120, anchor=tk.CENTER)

        # Grid layout
        self.tree.grid(row=0, column=0, sticky=tk.NSEW)
        sb_vert.grid(row=0, column=1, sticky=tk.NS)
        sb_horiz.grid(row=1, column=0, sticky=tk.EW)

        frm_tree.rowconfigure(0, weight=1)
        frm_tree.columnconfigure(0, weight=1)

        # Bindings
        self.tree.bind("<<TreeviewOpen>>", self._on_expand)
        self.tree.bind("<Double-Button-1>", self._on_double_click)

        # Statistics frame
        frm_stats = ttk.Frame(self.frm_main, style="App.TFrame")
        frm_stats.pack(side=tk.TOP, fill=tk.X, **paddings)

        self.lbl_stats = ttk.Label(frm_stats, text="")
        self.lbl_stats.pack(side=tk.LEFT)

        # Mandatory tests indicator (clickable)
        self.lbl_mandatory = ttk.Label(
            frm_stats,
            text="",
            foreground="red",
            cursor="hand2"
        )
        self.lbl_mandatory.pack(side=tk.RIGHT, padx=(20, 0))
        self.lbl_mandatory.bind("<Button-1>", self._on_show_missing_mandatory)

        # Buttons frame
        frm_buttons = ttk.Frame(self.frm_main, style="App.TFrame")
        frm_buttons.pack(side=tk.TOP, fill=tk.X, **paddings)

        self.btn_approve = ttk.Button(
            frm_buttons,
            text=_("Approve"),
            command=self._on_approve_workstation
        )
        self.btn_approve.pack(side=tk.LEFT, **paddings)

        self.btn_validate = ttk.Button(
            frm_buttons,
            text=_("Validate"),
            command=self._on_validate_result
        )
        self.btn_validate.pack(side=tk.LEFT, **paddings)

        self.btn_invalidate = ttk.Button(
            frm_buttons,
            text=_("Invalidate"),
            command=self._on_invalidate
        )
        self.btn_invalidate.pack(side=tk.LEFT, **paddings)

        ttk.Button(
            frm_buttons,
            text=_("Export"),
            command=self._on_export
        ).pack(side=tk.LEFT, **paddings)

        ttk.Button(
            frm_buttons,
            text=_("History"),
            command=self._on_show_history
        ).pack(side=tk.LEFT, **paddings)

        ttk.Button(
            frm_buttons,
            text=_("Close"),
            command=self.on_close
        ).pack(side=tk.RIGHT, **paddings)

    def on_open(self):
        """Entry point when opening the window."""
        self._check_user_permissions()
        self.calendarium.set_today()
        self._start_auto_refresh()
        self.deiconify()
        self.lift()

    def _start_auto_refresh(self):
        """Start auto-refresh every 30 seconds."""
        self._load_data()
        self._refresh_job = self.after(30000, self._start_auto_refresh)

    def _stop_auto_refresh(self):
        """Stop auto-refresh timer."""
        if self._refresh_job:
            self.after_cancel(self._refresh_job)
            self._refresh_job = None

    def _check_user_permissions(self):
        """Check user role and enable/disable validation controls."""
        try:
            user_role = self.engine.get_user_role()
            self.can_validate = user_role in (ROLE_ADMIN, ROLE_SUPERUSER)

            if self.can_validate:
                self.lbl_role.config(text=_("Validation enabled"), foreground="green")
            else:
                self.lbl_role.config(text=_("View-only mode"), foreground="orange")

            self._update_button_states()

        except Exception as e:
            self.can_validate = False
            self.lbl_role.config(text=_("View-only mode"), foreground="red")
            self.engine.on_log(
                "_check_user_permissions",
                e, type(e), sys.modules[__name__]
            )

    def _update_button_states(self):
        """Enable/disable buttons based on user role."""
        state = tk.NORMAL if self.can_validate else tk.DISABLED
        self.btn_approve.config(state=state)
        self.btn_validate.config(state=state)
        self.btn_invalidate.config(state=state)

    def _get_selected_date(self):
        """Get selected date from Calendarium."""
        try:
            value = self.calendarium.get_date()
            if value is False or value is None:
                return None
            return value
        except Exception:
            return None

    # =========================================================================
    # DATA LOADING
    # =========================================================================

    def _get_expanded_ws_ids(self):
        """Get list of currently expanded workstation IDs."""
        expanded = []
        for item_id in self.tree.get_children():
            if self.tree.item(item_id, "open"):
                row = self.dict_workstations.get(item_id)
                if row:
                    expanded.append(row["workstation_id"])
        return expanded

    def _expand_ws_ids(self, ws_ids):
        """Expand workstation nodes by their IDs."""
        for item_id, row in self.dict_workstations.items():
            if row["workstation_id"] in ws_ids:
                self.tree.item(item_id, open=True)
                # Trigger load of children
                self._on_expand(None, item_id)

    def _on_expand(self, evt, item_id=None):
        """Handle expand event - load results for workstation."""
        if item_id is None:
            item_id = self.tree.focus()
        if not item_id:
            return

        # Check if this is a workstation node
        tags = self.tree.item(item_id, "tags")
        if TAG_WORKSTATION not in tags:
            return

        row = self.dict_workstations.get(item_id)
        if not row:
            return

        ws_id = row["workstation_id"]

        # Already loaded?
        if ws_id in self.loaded_ws:
            return

        # Remove dummy child
        for child in self.tree.get_children(item_id):
            self.tree.delete(child)

        # Load results
        self._load_results_for_workstation(item_id, ws_id)
        self.loaded_ws.add(ws_id)

    def _load_data(self, preserve_expansion=False):
        """Load workstation summary for selected date (threaded)."""
        selected_date = self._get_selected_date()
        if selected_date is None:
            messagebox.showwarning(_("Validation"), _("Please select a valid date."))
            return

        # Save expanded state if requested
        expanded_ws_ids = self._get_expanded_ws_ids() if preserve_expansion else []

        self.selected_date = selected_date
        self.tree.delete(*self.tree.get_children())
        self.dict_workstations.clear()
        self.dict_results.clear()
        self.loaded_ws.clear()

        # Show loading state
        self.lbl_stats.config(text=_("Loading..."))
        self.btn_load.config(state=tk.DISABLED)

        # Run query in background thread
        def fetch_data():
            try:
                sql = """
                    SELECT
                        w.workstation_id,
                        w.description AS workstation_name,
                        e.description AS equipment_name,
                        COUNT(r.result_id) AS total_results,
                        SUM(CASE WHEN r.validated = 1 THEN 1 ELSE 0 END) AS validated_count,
                        SUM(CASE WHEN r.validated = 0 THEN 1 ELSE 0 END) AS pending_count,
                        SUM(CASE
                            WHEN r.validated = 0 AND b.sd > 0
                                 AND ABS(r.result - b.target) > (b.sd * 3)
                            THEN 1 ELSE 0
                        END) AS problem_count,
                        da.approval_id,
                        da.approved_by,
                        da.approved_at,
                        u.first_name,
                        u.last_name
                    FROM workstations w
                    INNER JOIN equipments e ON w.equipment_id = e.equipment_id
                    INNER JOIN sections s ON w.section_id = s.section_id
                    LEFT JOIN results r ON r.workstation_id = w.workstation_id
                        AND DATE(r.received) = ?
                        AND r.status = 1
                        AND r.is_delete = 0
                    LEFT JOIN batches b ON r.batch_id = b.batch_id
                        AND b.lab_id = ?
                    LEFT JOIN daily_approvals da ON da.workstation_id = w.workstation_id
                        AND da.approval_date = ?
                    LEFT JOIN users u ON da.approved_by = u.user_id
                    WHERE w.status = 1
                        AND s.lab_id = ?
                    GROUP BY w.workstation_id, w.description, e.description,
                             da.approval_id, da.approved_by, da.approved_at,
                             u.first_name, u.last_name
                    HAVING total_results > 0
                    ORDER BY w.description
                """

                lab_id = self.engine.current_ids.get("lab_id")
                args = (selected_date.isoformat(), lab_id, selected_date.isoformat(), lab_id)
                rows = self.engine.read(True, sql, args)

                # Update UI from main thread
                self.after(0, lambda: self._populate_tree(rows, expanded_ws_ids))

            except Exception as e:
                self.engine.on_log(
                    "_load_data",
                    e, type(e), sys.modules[__name__]
                )
                self.after(0, lambda: self._on_load_error(e))

        thread = threading.Thread(target=fetch_data, daemon=True)
        thread.start()

    def _populate_tree(self, rows, expanded_ws_ids):
        """Populate tree with fetched data (called from main thread)."""
        try:
            if rows is None:
                rows = []

            total_ws = 0
            approved_ws = 0

            for row in rows:
                total_ws += 1
                self._insert_workstation_node(row)
                if row["approval_id"]:
                    approved_ws += 1

            self.lbl_stats.config(
                text=f"Workstations: {total_ws}  |  Approved: {approved_ws}  |  Pending: {total_ws - approved_ws}"
            )

            # Check mandatory tests
            self._update_mandatory_indicator()

            # Restore expanded state
            if expanded_ws_ids:
                self._expand_ws_ids(expanded_ws_ids)

        finally:
            self.btn_load.config(state=tk.NORMAL)

    def _on_load_error(self, error):
        """Handle load error (called from main thread)."""
        self.btn_load.config(state=tk.NORMAL)
        self.lbl_stats.config(text="")
        messagebox.showerror(_("Error"), f"{_('Failed to load data:')}\n{error}")

    def _insert_workstation_node(self, row):
        """Insert a workstation as parent node."""
        ws_id = row["workstation_id"]
        ws_name = row["workstation_name"]
        eq_name = row["equipment_name"]
        total = row["total_results"] or 0
        validated = row["validated_count"] or 0
        pending = row["pending_count"] or 0
        problems = row["problem_count"] or 0
        approval_id = row["approval_id"]

        # Status text and color
        if approval_id:
            approved_by = f"{row['first_name'] or ''} {row['last_name'] or ''}".strip()
            status_text = f"✓ {approved_by}"
            color = self.engine.get_rgb(200, 255, 200)  # Green
        elif problems > 0:
            status_text = _("Problems")
            color = self.engine.get_rgb(255, 160, 160)  # Red
        elif pending > 0:
            status_text = _("Pending")
            color = self.engine.get_rgb(255, 255, 180)  # Yellow
        else:
            status_text = _("All validated")
            color = self.engine.get_rgb(200, 255, 200)  # Green

        counts_text = f"{total} tot / {pending} pend"
        problems_text = f"{problems} prob" if problems > 0 else ""

        # time column empty for workstations
        values = (eq_name, "", counts_text, problems_text, status_text)
        tags = (TAG_WORKSTATION, color)

        # Insert with dummy child so it's expandable
        item_id = self.tree.insert(
            "", tk.END,
            text=ws_name,
            values=values,
            tags=tags,
            open=False
        )

        # Add dummy child for expand arrow
        self.tree.insert(item_id, tk.END, text="Loading...")

        # Configure color
        self.tree.tag_configure(color, background=color)

        # Store data
        self.dict_workstations[item_id] = row

    def _load_results_for_workstation(self, parent_id, ws_id):
        """Load individual results as children of workstation node."""
        try:
            sql = """
                SELECT
                    r.result_id,
                    r.batch_id,
                    r.workstation_id,
                    r.result,
                    r.received,
                    r.validated,
                    r.validated_by,
                    b.target,
                    b.sd,
                    b.lot_number,
                    b.description AS level,
                    t.description AS test_description,
                    s.sample,
                    u.first_name AS validated_first_name,
                    u.last_name AS validated_last_name
                FROM results r
                INNER JOIN batches b ON r.batch_id = b.batch_id
                INNER JOIN test_methods tm ON b.test_method_id = tm.test_method_id
                INNER JOIN tests t ON tm.test_id = t.test_id
                INNER JOIN samples s ON tm.sample_id = s.sample_id
                LEFT JOIN users u ON r.validated_by = u.user_id
                WHERE r.workstation_id = ?
                  AND DATE(r.received) = ?
                  AND r.status = 1
                  AND r.is_delete = 0
                  AND b.lab_id = ?
                ORDER BY t.description, r.received
            """

            lab_id = self.engine.current_ids.get("lab_id")
            args = (ws_id, self.selected_date.isoformat(), lab_id)
            rows = self.engine.read(True, sql, args)

            if rows is None:
                rows = []

            for row in rows:
                self._insert_result_node(parent_id, row)

        except Exception as e:
            self.engine.on_log(
                "_load_results_for_workstation",
                e, type(e), sys.modules[__name__]
            )

    def _insert_result_node(self, parent_id, row):
        """Insert a result as child node under workstation."""
        result_val = float(row["result"])
        target = float(row["target"])
        sd = float(row["sd"])
        validated = row["validated"]
        lot_number = row.get("lot_number", "")
        level = row.get("level", "")
        test_desc = row["test_description"]
        sample = row["sample"]

        # Format time
        received = row["received"]
        if isinstance(received, datetime):
            time_str = received.strftime("%H:%M")
        else:
            time_str = str(received)[:5] if received else ""

        # Calculate Z-score
        if sd > 0:
            zscore = (result_val - target) / sd
            zscore_str = f"{zscore:+.2f} SD"
        else:
            zscore = 0
            zscore_str = "-"

        # Status
        if validated == 1:
            validated_by = f"{row.get('validated_first_name') or ''} {row.get('validated_last_name') or ''}".strip()
            status_text = f"✓ {validated_by}" if validated_by else "✓"
            color = self.engine.get_rgb(200, 255, 200)  # Green
        elif abs(zscore) >= 3:
            status_text = "⚠ >3SD"
            color = self.engine.get_rgb(255, 160, 160)  # Red
        elif abs(zscore) >= 2:
            status_text = "Pending"
            color = self.engine.get_rgb(255, 255, 180)  # Yellow
        else:
            status_text = "Pending"
            color = None

        test_name = f"{test_desc}-{sample}"
        batch_info = f"{lot_number} {level}".strip()
        result_str = f"{result_val:.2f}"

        values = (batch_info, time_str, result_str, zscore_str, status_text)
        tags = [TAG_RESULT]
        if color:
            tags.append(color)

        item_id = self.tree.insert(
            parent_id, tk.END,
            text=f"  {test_name}",
            values=values,
            tags=tuple(tags)
        )

        if color:
            self.tree.tag_configure(color, background=color)

        # Store data
        self.dict_results[item_id] = row

    # =========================================================================
    # ACTIONS
    # =========================================================================

    def _get_selected_item(self):
        """Get selected item and its type."""
        selection = self.tree.selection()
        if not selection:
            return None, None, None

        item_id = selection[0]
        tags = self.tree.item(item_id, "tags")

        if TAG_WORKSTATION in tags:
            return item_id, TAG_WORKSTATION, self.dict_workstations.get(item_id)
        elif TAG_RESULT in tags:
            return item_id, TAG_RESULT, self.dict_results.get(item_id)

        return None, None, None

    def _on_double_click(self, evt):
        """Handle double-click - validate single result."""
        if not self.can_validate:
            return

        item_id = self.tree.identify_row(evt.y)
        if not item_id:
            return

        tags = self.tree.item(item_id, "tags")
        if TAG_RESULT not in tags:
            return

        row = self.dict_results.get(item_id)
        if not row:
            return

        if row["validated"] == 1:
            messagebox.showinfo(_("Validation"), _("This result is already validated."))
            return

        self._validate_result(row["result_id"], item_id)

    def _on_approve_workstation(self):
        """Approve selected workstation."""
        if not self.can_validate:
            messagebox.showinfo(_("Validation"), _("You don't have permission to approve."))
            return

        item_id, item_type, row = self._get_selected_item()

        if item_type != TAG_WORKSTATION:
            messagebox.showinfo(_("Approve"), _("Please select a workstation."))
            return

        # Check if already approved
        if row["approval_id"]:
            messagebox.showinfo(_("Approve"), _("This workstation is already approved for today."))
            return

        # Check for problems
        if (row["problem_count"] or 0) > 0:
            if not messagebox.askyesno(
                _("Warning"),
                _("This workstation has {0} result(s) beyond ±3SD.\n\n"
                "Are you sure you want to approve it anyway?").format(row['problem_count'])
            ):
                return

        ws_name = row["workstation_name"]
        pending = row["pending_count"] or 0

        msg = _("Approve workstation '{0}'?").format(ws_name)
        if pending > 0:
            msg += _("\n\nThis will also validate {0} pending result(s).").format(pending)

        if not messagebox.askyesno(_("Confirm Approval"), msg):
            return

        try:
            user_id = self.engine.log_user.get("user_id")
            ws_id = row["workstation_id"]
            lab_id = self.engine.current_ids.get("lab_id")

            # 1. Validate all pending results for this workstation
            if pending > 0:
                sql_validate = """
                    UPDATE results r
                    INNER JOIN batches b ON r.batch_id = b.batch_id
                    SET r.validated = 1,
                        r.validated_by = ?,
                        r.validated_at = NOW()
                    WHERE r.workstation_id = ?
                      AND DATE(r.received) = ?
                      AND r.validated = 0
                      AND r.status = 1
                      AND r.is_delete = 0
                      AND b.lab_id = ?
                """
                self.engine.write(sql_validate, (user_id, ws_id, self.selected_date.isoformat(), lab_id))

            # 2. Insert approval record
            sql_approve = """
                INSERT INTO daily_approvals (approval_date, workstation_id, approved_by)
                VALUES (?, ?, ?)
            """
            self.engine.write(sql_approve, (self.selected_date.isoformat(), ws_id, user_id))

            messagebox.showinfo(_("Success"), _("Workstation '{0}' approved.").format(ws_name))
            self._load_data(preserve_expansion=True)

        except Exception as e:
            self.engine.on_log(
                "_on_approve_workstation",
                e, type(e), sys.modules[__name__]
            )
            messagebox.showerror(_("Error"), f"{_('Failed to approve:')}\n{e}")

    def _on_validate_result(self):
        """Validate selected result."""
        if not self.can_validate:
            messagebox.showinfo(_("Validation"), _("You don't have permission to validate."))
            return

        item_id, item_type, row = self._get_selected_item()

        if item_type != TAG_RESULT:
            messagebox.showinfo(_("Validate"), _("Please select a result."))
            return

        if row["validated"] == 1:
            messagebox.showinfo(_("Validation"), _("This result is already validated."))
            return

        self._validate_result(row["result_id"], item_id)

    def _validate_result(self, result_id, item_id):
        """Validate a single result."""
        try:
            user_id = self.engine.log_user.get("user_id")

            sql = """
                UPDATE results
                SET validated = 1,
                    validated_by = ?,
                    validated_at = NOW()
                WHERE result_id = ?
            """
            self.engine.write(sql, (user_id, result_id))

            # Update tree display
            row = self.dict_results[item_id]
            row["validated"] = 1
            ws_id = row["workstation_id"]

            color = self.engine.get_rgb(200, 255, 200)
            values = list(self.tree.item(item_id, "values"))
            values[4] = "✓"
            self.tree.item(item_id, values=values, tags=(TAG_RESULT, color))
            self.tree.tag_configure(color, background=color)

            # Check if all results are now validated → auto-approve WS
            self._check_auto_approve_workstation(ws_id, user_id)

            # Refresh parent workstation counts (preserve expansion)
            self._load_data(preserve_expansion=True)

        except Exception as e:
            self.engine.on_log(
                "_validate_result",
                e, type(e), sys.modules[__name__]
            )
            messagebox.showerror(_("Error"), f"{_('Failed to validate:')}\n{e}")

    def _check_auto_approve_workstation(self, ws_id, user_id):
        """Check if all results are validated and offer to approve workstation."""
        try:
            lab_id = self.engine.current_ids.get("lab_id")

            # Check if already approved
            sql_approved = """
                SELECT approval_id FROM daily_approvals
                WHERE workstation_id = ? AND approval_date = ?
            """
            approved = self.engine.read(True, sql_approved, (ws_id, self.selected_date.isoformat()))
            if approved:
                return  # Already approved

            # Count pending results
            sql_pending = """
                SELECT COUNT(*) AS pending
                FROM results r
                INNER JOIN batches b ON r.batch_id = b.batch_id
                WHERE r.workstation_id = ?
                  AND DATE(r.received) = ?
                  AND r.validated = 0
                  AND r.status = 1
                  AND r.is_delete = 0
                  AND b.lab_id = ?
            """
            result = self.engine.read(True, sql_pending, (ws_id, self.selected_date.isoformat(), lab_id))
            pending = result[0]["pending"] if result else 0

            if pending > 0:
                return  # Still has pending results

            # All validated! Offer to approve workstation
            if messagebox.askyesno(
                _("All Validated"),
                _("All results for this workstation are now validated.\n\n"
                "Approve the workstation?")
            ):
                sql_approve = """
                    INSERT INTO daily_approvals (approval_date, workstation_id, approved_by)
                    VALUES (?, ?, ?)
                """
                self.engine.write(sql_approve, (self.selected_date.isoformat(), ws_id, user_id))

        except Exception as e:
            self.engine.on_log(
                "_check_auto_approve_workstation",
                e, type(e), sys.modules[__name__]
            )

    def _on_invalidate(self):
        """Invalidate a validated result."""
        if not self.can_validate:
            messagebox.showinfo(_("Validation"), _("You don't have permission to invalidate."))
            return

        item_id, item_type, row = self._get_selected_item()

        if item_type != TAG_RESULT:
            messagebox.showinfo(_("Invalidate"), _("Please select a result."))
            return

        if row["validated"] == 0:
            messagebox.showinfo(_("Invalidate"), _("This result is not validated."))
            return

        ws_id = row["workstation_id"]

        # Check if workstation is approved
        ws_approved = self._is_workstation_approved(ws_id)

        if ws_approved:
            msg = _(
                "Invalidate result {0}?\n\n"
                "This workstation is approved.\n"
                "The approval will be revoked."
            ).format(row['result_id'])
        else:
            msg = _("Invalidate result {0}?").format(row['result_id'])

        if not messagebox.askyesno(_("Confirm"), msg):
            return

        try:
            # Invalidate the result
            sql = """
                UPDATE results
                SET validated = 0,
                    validated_by = NULL,
                    validated_at = NULL
                WHERE result_id = ?
            """
            self.engine.write(sql, (row["result_id"],))

            # Revoke workstation approval if it was approved
            if ws_approved:
                sql_revoke = """
                    DELETE FROM daily_approvals
                    WHERE workstation_id = ? AND approval_date = ?
                """
                self.engine.write(sql_revoke, (ws_id, self.selected_date.isoformat()))
                messagebox.showinfo(_("Success"), _("Result invalidated. Workstation approval revoked."))
            else:
                messagebox.showinfo(_("Success"), _("Result invalidated."))

            self._load_data(preserve_expansion=True)

        except Exception as e:
            self.engine.on_log(
                "_on_invalidate",
                e, type(e), sys.modules[__name__]
            )
            messagebox.showerror(_("Error"), f"{_('Failed to invalidate:')}\n{e}")

    def _is_workstation_approved(self, ws_id):
        """Check if workstation is approved for selected date."""
        try:
            sql = """
                SELECT approval_id FROM daily_approvals
                WHERE workstation_id = ? AND approval_date = ?
            """
            result = self.engine.read(True, sql, (ws_id, self.selected_date.isoformat()))
            return bool(result)
        except Exception:
            return False

    def _on_export(self):
        """Export daily validation data to Excel."""
        if not self.selected_date:
            messagebox.showinfo(_("Export"), _("Please load data first."))
            return

        try:
            self.engine.quick_data_analysis(self.selected_date, None)
        except Exception as e:
            self.engine.on_log(
                "_on_export",
                e, type(e), sys.modules[__name__]
            )
            messagebox.showerror(_("Error"), f"{_('Failed to export:')}\n{e}")

    # =========================================================================
    # VALIDATION HISTORY
    # =========================================================================

    def _on_show_history(self):
        """Show validation history for selected date."""
        if not self.selected_date:
            messagebox.showinfo(_("History"), _("Please load data first."))
            return

        history = self._get_validation_history()

        # Create popup window
        popup = tk.Toplevel(self)
        popup.title(f"Validation History - {self.selected_date}")
        popup.geometry("800x450")
        popup.transient(self)

        # Main frame with grid layout
        frm = ttk.Frame(popup, padding=10)
        frm.pack(fill=tk.BOTH, expand=True)
        frm.rowconfigure(1, weight=1)
        frm.columnconfigure(0, weight=1)

        # Info label (row 0)
        ttk.Label(
            frm,
            text=f"Validation actions for {self.selected_date}",
            font=("TkDefaultFont", 10, "bold")
        ).grid(row=0, column=0, sticky=tk.W, pady=(0, 10))

        # Treeview with scrollbar (row 1)
        frm_tree = ttk.Frame(frm)
        frm_tree.grid(row=1, column=0, sticky=tk.NSEW)
        frm_tree.rowconfigure(0, weight=1)
        frm_tree.columnconfigure(0, weight=1)

        sb = ttk.Scrollbar(frm_tree, orient=tk.VERTICAL)
        cols = ("time", "user", "test", "workstation", "action", "result_value")
        tree = ttk.Treeview(frm_tree, columns=cols, show="headings", yscrollcommand=sb.set, height=15)
        sb.config(command=tree.yview)

        tree.heading("time", text="Time", anchor=tk.CENTER)
        tree.heading("user", text="User", anchor=tk.W)
        tree.heading("test", text="Test", anchor=tk.W)
        tree.heading("workstation", text="Workstation", anchor=tk.W)
        tree.heading("action", text="Action", anchor=tk.CENTER)
        tree.heading("result_value", text="Value", anchor=tk.E)

        tree.column("time", width=60, anchor=tk.CENTER)
        tree.column("user", width=120, anchor=tk.W)
        tree.column("test", width=180, anchor=tk.W)
        tree.column("workstation", width=150, anchor=tk.W)
        tree.column("action", width=80, anchor=tk.CENTER)
        tree.column("result_value", width=80, anchor=tk.E)

        tree.grid(row=0, column=0, sticky=tk.NSEW)
        sb.grid(row=0, column=1, sticky=tk.NS)

        # Populate
        for row in history:
            action = "Validated" if row["validated"] == 1 else "Invalidated"
            color = self.engine.get_rgb(200, 255, 200) if row["validated"] == 1 else self.engine.get_rgb(255, 200, 200)

            values = (
                row["time_str"],
                row["user_name"],
                row["test_name"],
                row["workstation_name"],
                action,
                f"{row['result_value']:.2f}" if row["result_value"] else ""
            )
            item = tree.insert("", tk.END, values=values, tags=(color,))
            tree.tag_configure(color, background=color)

        # Stats (row 2)
        validated_count = sum(1 for r in history if r["validated"] == 1)
        invalidated_count = sum(1 for r in history if r["validated"] == 0)

        ttk.Label(
            frm,
            text=f"Total: {len(history)}  |  Validated: {validated_count}  |  Invalidated: {invalidated_count}"
        ).grid(row=2, column=0, sticky=tk.W, pady=(10, 0))

        # Buttons (row 3)
        frm_btn = ttk.Frame(frm)
        frm_btn.grid(row=3, column=0, sticky=tk.EW, pady=(10, 0))

        ttk.Button(
            frm_btn,
            text="Export",
            command=lambda: self._export_history(history)
        ).pack(side=tk.LEFT, padx=5)

        ttk.Button(
            frm_btn,
            text="Close",
            command=popup.destroy
        ).pack(side=tk.RIGHT, padx=5)

    def _get_validation_history(self):
        """Get validation history from audit_results for selected date."""
        try:
            lab_id = self.engine.current_ids.get("lab_id")

            sql = """
                SELECT
                    ar.log_time,
                    ar.validated,
                    ar.result AS result_value,
                    u.first_name,
                    u.last_name,
                    t.description AS test_description,
                    s.sample,
                    w.description AS workstation_name
                FROM audit_results ar
                INNER JOIN results r ON ar.result_id = r.result_id
                INNER JOIN batches b ON r.batch_id = b.batch_id
                INNER JOIN test_methods tm ON b.test_method_id = tm.test_method_id
                INNER JOIN tests t ON tm.test_id = t.test_id
                INNER JOIN samples s ON tm.sample_id = s.sample_id
                INNER JOIN workstations w ON r.workstation_id = w.workstation_id
                LEFT JOIN users u ON ar.validated_by = u.user_id
                WHERE DATE(ar.log_time) = ?
                  AND ar.validated_by IS NOT NULL
                  AND b.lab_id = ?
                ORDER BY ar.log_time DESC
            """

            rows = self.engine.read(True, sql, (self.selected_date.isoformat(), lab_id))

            if not rows:
                return []

            history = []
            for row in rows:
                log_time = row["log_time"]
                if isinstance(log_time, datetime):
                    time_str = log_time.strftime("%H:%M")
                else:
                    time_str = str(log_time)[:5] if log_time else ""

                user_name = f"{row['first_name'] or ''} {row['last_name'] or ''}".strip() or "Unknown"
                test_name = f"{row['test_description']}-{row['sample']}"

                history.append({
                    "time_str": time_str,
                    "user_name": user_name,
                    "test_name": test_name,
                    "workstation_name": row["workstation_name"],
                    "validated": row["validated"],
                    "result_value": row["result_value"]
                })

            return history

        except Exception as e:
            self.engine.on_log(
                "_get_validation_history",
                e, type(e), sys.modules[__name__]
            )
            return []

    def _export_history(self, history):
        """Export validation history to Excel."""
        if not history:
            messagebox.showinfo(_("Export"), _("No data to export."))
            return

        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
            import os

            wb = Workbook()
            ws = wb.active
            ws.title = "Validation History"

            # Header
            headers = ["Time", "User", "Test", "Workstation", "Action", "Value"]
            ws.append(headers)
            for cell in ws[1]:
                cell.font = Font(bold=True)

            # Data
            for row in history:
                action = "Validated" if row["validated"] == 1 else "Invalidated"
                ws.append([
                    row["time_str"],
                    row["user_name"],
                    row["test_name"],
                    row["workstation_name"],
                    action,
                    row["result_value"]
                ])

            # Column widths
            ws.column_dimensions['A'].width = 8
            ws.column_dimensions['B'].width = 20
            ws.column_dimensions['C'].width = 25
            ws.column_dimensions['D'].width = 20
            ws.column_dimensions['E'].width = 12
            ws.column_dimensions['F'].width = 10

            # Save
            filename = f"validation_history_{self.selected_date}.xlsx"
            filepath = os.path.join(os.path.dirname(os.path.dirname(__file__)), filename)
            wb.save(filepath)

            messagebox.showinfo(_("Export"), _("History exported to:\n{0}").format(filename))

            # Open file
            if sys.platform == "win32":
                os.startfile(filepath)
            else:
                import subprocess
                subprocess.run(["xdg-open", filepath], check=False)

        except Exception as e:
            self.engine.on_log(
                "_export_history",
                e, type(e), sys.modules[__name__]
            )
            messagebox.showerror(_("Error"), f"{_('Failed to export:')}\n{e}")

    # =========================================================================
    # MANDATORY TESTS
    # =========================================================================

    def _update_mandatory_indicator(self):
        """Update the mandatory tests indicator label."""
        missing = self._get_missing_mandatory()
        self.missing_mandatory = missing  # Store for click handler

        if missing:
            self.lbl_mandatory.config(
                text=f"⚠ Missing mandatory: {len(missing)} (click)",
                foreground="red"
            )
        else:
            self.lbl_mandatory.config(
                text="✓ All mandatory OK",
                foreground="green"
            )

    def _get_missing_mandatory(self):
        """Get list of mandatory tests not executed for selected date."""
        if not self.selected_date:
            return []

        try:
            # Get mandatory tests for current section
            mandatory = self.engine.get_mandatory()
            if not mandatory:
                return []

            # Get tests executed today
            lab_id = self.engine.current_ids.get("lab_id")
            sql = """
                SELECT DISTINCT t.description
                FROM results r
                INNER JOIN batches b ON r.batch_id = b.batch_id
                INNER JOIN test_methods tm ON b.test_method_id = tm.test_method_id
                INNER JOIN tests t ON tm.test_id = t.test_id
                WHERE DATE(r.received) = ?
                  AND r.status = 1
                  AND r.is_delete = 0
                  AND b.lab_id = ?
            """
            rows = self.engine.read(True, sql, (self.selected_date.isoformat(), lab_id))

            executed = set()
            if rows:
                executed = {row["description"] for row in rows}

            # Find missing
            missing = [t for t in mandatory if t not in executed]
            return missing

        except Exception as e:
            self.engine.on_log(
                "_get_missing_mandatory",
                e, type(e), sys.modules[__name__]
            )
            return []

    def _on_show_missing_mandatory(self, evt=None):
        """Show popup with list of missing mandatory tests."""
        if not hasattr(self, 'missing_mandatory') or not self.missing_mandatory:
            messagebox.showinfo(_("Mandatory Tests"), _("All mandatory tests have been executed."))
            return

        missing_list = "\n".join(f"  • {t}" for t in self.missing_mandatory)
        messagebox.showwarning(
            _("Missing Mandatory Tests"),
            _("The following mandatory tests have not been executed:\n\n{0}").format(missing_list)
        )

    def on_close(self, evt=None):
        """Close the window."""
        self._stop_auto_refresh()
        self.engine.dict_instances.pop(self.winfo_name(), None)
        self.destroy()
