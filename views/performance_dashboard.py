# -*- coding: utf-8 -*-
#-----------------------------------------------------------------------------
# project:  biovarase
# authors:  1966bc
# mailto:   giuseppecostanzi@gmail.com
# modify:   ver MMXXVI
#-----------------------------------------------------------------------------
"""
Performance Dashboard - Monitor test method performance metrics.

Shows aggregated QC metrics per test method for a selected date range:
- Westgard violation rate
- Warning rate
- Notes/actions rate
- Observed CV% and Bias%
"""

import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from datetime import datetime, timedelta
from decimal import Decimal

from i18n import _
from views.parent_view import ParentView
from calendarium import Calendarium


# SQL to get results with test method info
SQL_RESULTS = """
    SELECT
        tm.test_method_id,
        t.description AS test_name,
        s.description AS sample_type,
        w.description AS workstation,
        w.workstation_id,
        r.result_id,
        r.result,
        r.received,
        b.batch_id,
        b.target,
        b.sd,
        (SELECT COUNT(*) FROM notes n WHERE n.result_id = r.result_id AND n.status = 1) AS note_count
    FROM results r
    INNER JOIN batches b ON r.batch_id = b.batch_id
    INNER JOIN test_methods tm ON b.test_method_id = tm.test_method_id
    INNER JOIN tests t ON tm.test_id = t.test_id
    INNER JOIN samples s ON tm.sample_id = s.sample_id
    INNER JOIN workstations w ON r.workstation_id = w.workstation_id
    INNER JOIN organizations section ON section.org_id = tm.org_id
    WHERE DATE(r.received) BETWEEN ? AND ?
      AND section.parent_id = ?
      AND section.org_type = 'section'
      AND r.status = 1
      AND r.is_delete = 0
      AND tm.status = 1
      AND t.status = 1
    ORDER BY t.description, w.description, r.received
"""

# SQL to get action breakdown for a test method
SQL_ACTIONS = """
    SELECT
        a.description AS action_name,
        COUNT(*) AS action_count
    FROM notes n
    INNER JOIN actions a ON n.action_id = a.action_id
    INNER JOIN results r ON n.result_id = r.result_id
    INNER JOIN batches b ON r.batch_id = b.batch_id
    WHERE b.test_method_id = ?
      AND r.workstation_id = ?
      AND DATE(r.received) BETWEEN ? AND ?
      AND n.status = 1
      AND r.status = 1
      AND r.is_delete = 0
    GROUP BY a.action_id, a.description
    ORDER BY action_count DESC
"""


class UI(ParentView):
    """
    Performance Dashboard window.

    Displays aggregated QC metrics per test method for monitoring
    analytical performance over time.
    """

    def __init__(self, parent):
        super().__init__(parent, name="performance_dashboard")
        if self._reusing:
            self.on_open()
            return

        # Note: transient() removed - breaks resize on Windows
        self.resizable(True, True)
        self.geometry("1000x600")
        self.minsize(800, 400)

        # Hotkeys
        self.bind("<Alt-r>", self._on_refresh)
        self.bind("<Alt-e>", self._on_export)
        self.bind("<Alt-c>", self.on_cancel)

        # Data storage
        self.dict_items = {}      # tree item_id -> aggregated data
        self.aggregated = {}      # (tm_id, ws_id) -> metrics dict
        self.current_from = None
        self.current_to = None

        # Sort state
        self.sort_column = None
        self.sort_reverse = False

        self._build_ui()
        self.show()

    def _build_ui(self):
        """Build the UI components."""
        self.title(_("Performance Dashboard"))
        padd = {"padx": 5, "pady": 5}

        # Main container
        self.frm_main = ttk.Frame(self, style="App.TFrame", padding=8)
        self.frm_main.pack(fill=tk.BOTH, expand=True)

        # Top frame - Date selectors and buttons
        frm_top = ttk.Frame(self.frm_main, style="App.TFrame")
        frm_top.pack(fill=tk.X, **padd)

        # Get background color for Calendarium
        try:
            bg = self.engine.get_rgb(240, 240, 237)
        except Exception:
            bg = "#f0f0ed"

        # From date
        ttk.Label(frm_top, text=_("From:"), style="App.TLabel").pack(side=tk.LEFT, **padd)
        self.cal_from = Calendarium(frm_top, "", base_bg_color=bg)
        self.cal_from.pack(side=tk.LEFT, **padd)

        # To date
        ttk.Label(frm_top, text=_("To:"), style="App.TLabel").pack(side=tk.LEFT, **padd)
        self.cal_to = Calendarium(frm_top, "", base_bg_color=bg)
        self.cal_to.pack(side=tk.LEFT, **padd)

        # Buttons
        ttk.Button(
            frm_top,
            text=_("Refresh"),
            command=self._on_refresh,
        ).pack(side=tk.LEFT, **padd)

        ttk.Button(
            frm_top,
            text=_("Export"),
            command=self._on_export,
        ).pack(side=tk.LEFT, **padd)

        # Color filter
        ttk.Label(frm_top, text=_("Filter:"), style="App.TLabel").pack(side=tk.LEFT, padx=(15, 5))
        self.filter_var = tk.StringVar(value="all")
        self.cbx_filter = ttk.Combobox(
            frm_top,
            textvariable=self.filter_var,
            state="readonly",
            width=12,
            values=[
                _("All"),
                _("Critical"),    # Red
                _("Warning"),     # Yellow
                _("Good"),        # Green
            ],
        )
        self.cbx_filter.current(0)
        self.cbx_filter.pack(side=tk.LEFT, **padd)
        self.cbx_filter.bind("<<ComboboxSelected>>", self._on_filter_changed)

        ttk.Button(
            frm_top,
            text=_("Close"),
            command=self.on_cancel,
        ).pack(side=tk.RIGHT, **padd)

        # Status label
        self.status_var = tk.StringVar(value="")
        ttk.Label(frm_top, textvariable=self.status_var, style="App.TLabel").pack(side=tk.RIGHT, **padd)

        # PanedWindow for tree and details
        self.pw = tk.PanedWindow(self.frm_main, orient=tk.VERTICAL, sashwidth=6)
        self.pw.pack(fill=tk.BOTH, expand=True, **padd)

        # Main Treeview frame
        frm_tree = ttk.Frame(self.pw, style="App.TFrame")
        self.pw.add(frm_tree, minsize=200, stretch="always")

        # Scrollbars
        sb_vert = ttk.Scrollbar(frm_tree, orient=tk.VERTICAL)
        sb_vert.pack(side=tk.RIGHT, fill=tk.Y)
        sb_horiz = ttk.Scrollbar(frm_tree, orient=tk.HORIZONTAL)
        sb_horiz.pack(side=tk.BOTTOM, fill=tk.X)

        # Treeview columns
        cols = ("test", "workstation", "total", "viol_pct", "warn_pct", "note_pct", "cv_pct", "bias_pct")
        self.tree = ttk.Treeview(
            frm_tree,
            columns=cols,
            show="headings",
            yscrollcommand=sb_vert.set,
            xscrollcommand=sb_horiz.set,
        )
        sb_vert.config(command=self.tree.yview)
        sb_horiz.config(command=self.tree.xview)

        # Column definitions
        col_defs = [
            ("test", _("Test"), 200, tk.W),
            ("workstation", _("Workstation"), 100, tk.W),
            ("total", _("Total"), 60, tk.CENTER),
            ("viol_pct", _("Viol%"), 70, tk.CENTER),
            ("warn_pct", _("Warn%"), 70, tk.CENTER),
            ("note_pct", _("Note%"), 70, tk.CENTER),
            ("cv_pct", _("CV%"), 70, tk.CENTER),
            ("bias_pct", _("Bias%"), 80, tk.CENTER),
        ]

        for col_id, heading, width, anchor in col_defs:
            self.tree.column(col_id, width=width, minwidth=50, anchor=anchor, stretch=True)
            self.tree.heading(
                col_id,
                text=heading,
                anchor=anchor,
                command=lambda c=col_id: self._on_column_click(c),
            )

        self.tree.pack(fill=tk.BOTH, expand=True)

        # Color tags
        self.tree.tag_configure("high_violation", background="#ffcccc")     # Light red
        self.tree.tag_configure("medium_violation", background="#fff2cc")   # Light yellow
        self.tree.tag_configure("good", background="#ccffcc")               # Light green

        # Bind selection
        self.tree.bind("<<TreeviewSelect>>", self._on_tree_select)

        # Details panel
        frm_details = ttk.LabelFrame(self.pw, text=_("Actions Breakdown"))
        self.pw.add(frm_details, minsize=100, stretch="never")

        # Details listbox with scrollbar
        sb_details = ttk.Scrollbar(frm_details, orient=tk.VERTICAL)
        sb_details.pack(side=tk.RIGHT, fill=tk.Y)

        self.lst_details = tk.Listbox(frm_details, yscrollcommand=sb_details.set, height=5)
        self.lst_details.pack(fill=tk.BOTH, expand=True, padx=5, pady=5)
        sb_details.config(command=self.lst_details.yview)

        # Sizegrip for resize handle
        ttk.Sizegrip(self).place(relx=1.0, rely=1.0, anchor="se")

    def on_open(self):
        """Initialize with default date range (last 30 days)."""
        # Set default date range
        today = datetime.now().date()
        from_date = today - timedelta(days=30)

        self.cal_from.set_date(from_date)
        self.cal_to.set_date(today)

        self._on_refresh()

    def _on_refresh(self, evt=None):
        """Load and aggregate data for selected date range."""
        # Validate dates
        from_date = self.cal_from.get_date()
        to_date = self.cal_to.get_date()

        if from_date is None or to_date is None:
            messagebox.showwarning(
                self.engine.app_title,
                _("Invalid date."),
                parent=self,
            )
            return

        if from_date > to_date:
            messagebox.showwarning(
                self.engine.app_title,
                _("From date must be before To date."),
                parent=self,
            )
            return

        self.current_from = from_date
        self.current_to = to_date

        self.engine.busy(self)
        try:
            self._load_data(from_date, to_date)
        finally:
            self.engine.not_busy(self)

    def _load_data(self, from_date, to_date):
        """Query results and aggregate by test_method + workstation."""
        lab_id = self.engine.current_ids.get("lab_id")

        rows = self.engine.read(
            True,
            SQL_RESULTS,
            (from_date.isoformat(), to_date.isoformat(), lab_id),
        ) or []

        if not rows:
            self.engine.clear_treeview(self.tree)
            self.dict_items.clear()
            self.aggregated.clear()
            self.lst_details.delete(0, tk.END)
            self.status_var.set(_("No data for selected date range"))
            return

        # Aggregate results
        self.aggregated = self._aggregate_results(rows)

        # Populate tree
        self._populate_tree()

        # Update status
        count = len(self.aggregated)
        self.status_var.set(
            f"{_('Test Methods')}: {count} | "
            f"{from_date.strftime('%d/%m/%Y')} - {to_date.strftime('%d/%m/%Y')}"
        )

    def _aggregate_results(self, rows):
        """
        Group results by (test_method_id, workstation_id) and calculate metrics.

        Returns dict: {(tm_id, ws_id): {metrics...}}
        """
        aggregated = {}

        for row in rows:
            key = (row["test_method_id"], row["workstation_id"])

            if key not in aggregated:
                aggregated[key] = {
                    "test_name": row["test_name"],
                    "sample": row["sample_type"],
                    "workstation": row["workstation"],
                    "workstation_id": row["workstation_id"],
                    "test_method_id": row["test_method_id"],
                    "total": 0,
                    "violations": 0,
                    "warnings": 0,
                    "with_notes": 0,
                    "results": [],
                    "target": None,
                    "sd": None,
                }

            data = aggregated[key]
            data["total"] += 1

            # Collect results for CV/Bias calculation
            result_val = row["result"]
            if result_val is not None:
                try:
                    data["results"].append(float(result_val))
                except (TypeError, ValueError):
                    pass

            # Store target and sd (use latest batch values)
            target = row["target"]
            sd = row["sd"]
            if target is not None and sd is not None:
                try:
                    data["target"] = float(target)
                    data["sd"] = float(sd)
                except (TypeError, ValueError):
                    pass

            # Count notes
            note_count = row.get("note_count", 0) or 0
            if note_count > 0:
                data["with_notes"] += 1

            # Calculate Westgard (simplified z-score approach)
            self._evaluate_result(row, data)

        return aggregated

    def _evaluate_result(self, row, data):
        """
        Evaluate a single result for Westgard violations.

        Simplified approach using z-score:
        - |z| >= 3: Violation (1:3S equivalent)
        - 2 <= |z| < 3: Warning (1:2S equivalent)
        """
        result = row["result"]
        target = row["target"]
        sd = row["sd"]

        if result is None or target is None or sd is None:
            return

        try:
            result_f = float(result)
            target_f = float(target)
            sd_f = float(sd)
        except (TypeError, ValueError):
            return

        if sd_f <= 0:
            return

        zscore = abs((result_f - target_f) / sd_f)

        if zscore >= 3:
            data["violations"] += 1
        elif zscore >= 2:
            data["warnings"] += 1

    def _calculate_metrics(self, data):
        """Calculate CV% and Bias% for aggregated data."""
        results = data.get("results", [])
        target = data.get("target")

        if len(results) < 2 or target is None:
            return None, None

        try:
            avg = self.engine.get_mean(results)
            cv = self.engine.get_cv(results)
            bias = self.engine.get_bias(avg, target)
            return cv, bias
        except Exception:
            return None, None

    def _populate_tree(self):
        """Populate treeview with aggregated data."""
        self.engine.clear_treeview(self.tree)
        self.dict_items.clear()
        self.lst_details.delete(0, tk.END)

        # Get current filter
        filter_idx = self.cbx_filter.current()
        # 0 = All, 1 = Critical (red), 2 = Warning (yellow), 3 = Good (green)

        # Prepare data for sorting
        items = []
        for key, data in self.aggregated.items():
            cv, bias = self._calculate_metrics(data)

            total = data["total"]
            viol_pct = (data["violations"] / total * 100) if total > 0 else 0
            warn_pct = (data["warnings"] / total * 100) if total > 0 else 0
            note_pct = (data["with_notes"] / total * 100) if total > 0 else 0

            # Store calculated values for sorting
            data["viol_pct"] = viol_pct
            data["warn_pct"] = warn_pct
            data["note_pct"] = note_pct
            data["cv_pct"] = cv
            data["bias_pct"] = bias

            items.append((key, data))

        # Sort if column selected
        if self.sort_column:
            items = self._sort_items(items)

        # Insert into tree (with filter)
        shown_count = 0
        for key, data in items:
            tags = self._get_row_tag(data["viol_pct"], data["warn_pct"])

            # Apply filter
            if filter_idx == 1 and tags[0] != "high_violation":
                continue
            elif filter_idx == 2 and tags[0] != "medium_violation":
                continue
            elif filter_idx == 3 and tags[0] != "good":
                continue

            cv = data["cv_pct"]
            bias = data["bias_pct"]

            values = (
                f"{data['test_name']} - {data['sample']}",
                data["workstation"],
                data["total"],
                f"{data['viol_pct']:.1f}%",
                f"{data['warn_pct']:.1f}%",
                f"{data['note_pct']:.1f}%",
                f"{cv:.2f}%" if cv is not None else "-",
                f"{bias:+.2f}%" if bias is not None else "-",
            )

            item_id = self.tree.insert("", tk.END, values=values, tags=tags)
            self.dict_items[item_id] = data
            shown_count += 1

        # Update status with filter info
        total_count = len(self.aggregated)
        if self.current_from and self.current_to:
            period = f"{self.current_from.strftime('%d/%m/%Y')} - {self.current_to.strftime('%d/%m/%Y')}"
            if filter_idx > 0:
                self.status_var.set(f"{_('Showing')}: {shown_count}/{total_count} | {period}")
            else:
                self.status_var.set(f"{_('Test Methods')}: {total_count} | {period}")

    def _get_row_tag(self, viol_pct, warn_pct):
        """Return color tag based on violation/warning rates."""
        if viol_pct >= 5.0:
            return ("high_violation",)
        elif viol_pct >= 2.0 or warn_pct >= 10.0:
            return ("medium_violation",)
        else:
            return ("good",)

    def _sort_items(self, items):
        """Sort items by current sort column."""
        col = self.sort_column

        # Map column to data key
        col_map = {
            "test": lambda d: f"{d['test_name']} - {d['sample']}".lower(),
            "workstation": lambda d: d["workstation"].lower(),
            "total": lambda d: d["total"],
            "viol_pct": lambda d: d["viol_pct"],
            "warn_pct": lambda d: d["warn_pct"],
            "note_pct": lambda d: d["note_pct"],
            "cv_pct": lambda d: d["cv_pct"] if d["cv_pct"] is not None else -999,
            "bias_pct": lambda d: d["bias_pct"] if d["bias_pct"] is not None else -999,
        }

        key_func = col_map.get(col)
        if key_func:
            items.sort(key=lambda x: key_func(x[1]), reverse=self.sort_reverse)

        return items

    def _on_column_click(self, col):
        """Handle column header click for sorting."""
        if self.sort_column == col:
            self.sort_reverse = not self.sort_reverse
        else:
            self.sort_column = col
            self.sort_reverse = False

        self._populate_tree()

    def _on_filter_changed(self, evt=None):
        """Handle filter combobox selection change."""
        self._populate_tree()

    def _on_tree_select(self, evt=None):
        """Show action breakdown for selected test method."""
        selection = self.tree.selection()
        if not selection:
            self.lst_details.delete(0, tk.END)
            return

        item_id = selection[0]
        data = self.dict_items.get(item_id)
        if not data:
            return

        self._show_action_breakdown(data)

    def _show_action_breakdown(self, data):
        """Query and display action breakdown for selected test method."""
        self.lst_details.delete(0, tk.END)

        if self.current_from is None or self.current_to is None:
            return

        rows = self.engine.read(
            True,
            SQL_ACTIONS,
            (
                data["test_method_id"],
                data["workstation_id"],
                self.current_from.isoformat(),
                self.current_to.isoformat(),
            ),
        ) or []

        if not rows:
            self.lst_details.insert(tk.END, _("No actions recorded"))
            return

        # Calculate total for percentages
        total = sum(r["action_count"] for r in rows)

        for row in rows:
            action = _(row["action_name"])
            count = row["action_count"]
            pct = (count / total * 100) if total > 0 else 0
            self.lst_details.insert(tk.END, f"{action}: {count} ({pct:.0f}%)")

    def _on_export(self, evt=None):
        """Show export preview dialog."""
        if not self.dict_items:
            messagebox.showwarning(
                self.engine.app_title,
                _("No data to export."),
                parent=self,
            )
            return

        # Create preview window
        preview = tk.Toplevel(self)
        preview.title(_("Export Preview"))
        preview.transient(self)
        preview.geometry("800x500")
        preview.minsize(600, 400)

        # Header frame
        frm_header = ttk.Frame(preview, style="App.TFrame")
        frm_header.pack(fill=tk.X, padx=10, pady=5)

        period_text = ""
        if self.current_from and self.current_to:
            period_text = f"{self.current_from.strftime('%d/%m/%Y')} - {self.current_to.strftime('%d/%m/%Y')}"

        ttk.Label(
            frm_header,
            text=f"{_('Performance Dashboard')} - {period_text}",
            style="App.TLabel",
            font=("TkDefaultFont", 12, "bold"),
        ).pack(side=tk.LEFT)

        # Text widget with scrollbar
        frm_text = ttk.Frame(preview, style="App.TFrame")
        frm_text.pack(fill=tk.BOTH, expand=True, padx=10, pady=5)

        sb = ttk.Scrollbar(frm_text, orient=tk.VERTICAL)
        sb.pack(side=tk.RIGHT, fill=tk.Y)

        txt = tk.Text(frm_text, yscrollcommand=sb.set, font=("Courier", 10), wrap=tk.NONE)
        txt.pack(fill=tk.BOTH, expand=True)
        sb.config(command=txt.yview)

        # Configure tags for colors
        txt.tag_configure("header", font=("Courier", 10, "bold"), background="#d0d0d0")
        txt.tag_configure("high_violation", background="#ffcccc")
        txt.tag_configure("medium_violation", background="#fff2cc")
        txt.tag_configure("good", background="#ccffcc")

        # Build preview content
        headers = f"{'Test':<35} {'WS':<10} {'Tot':>5} {'Viol%':>7} {'Warn%':>7} {'Note%':>7} {'CV%':>7} {'Bias%':>8}\n"
        separator = "-" * 95 + "\n"

        txt.insert(tk.END, headers, "header")
        txt.insert(tk.END, separator)

        for item_id in self.tree.get_children():
            data = self.dict_items.get(item_id)
            if not data:
                continue

            values = self.tree.item(item_id, "values")
            tag = self._get_row_tag(data["viol_pct"], data["warn_pct"])[0]

            line = f"{values[0]:<35} {values[1]:<10} {values[2]:>5} {values[3]:>7} {values[4]:>7} {values[5]:>7} {values[6]:>7} {values[7]:>8}\n"
            txt.insert(tk.END, line, tag)

        txt.config(state=tk.DISABLED)

        # Horizontal scrollbar
        sb_h = ttk.Scrollbar(frm_text, orient=tk.HORIZONTAL, command=txt.xview)
        sb_h.pack(side=tk.BOTTOM, fill=tk.X)
        txt.config(xscrollcommand=sb_h.set)

        # Button frame
        frm_buttons = ttk.Frame(preview, style="App.TFrame")
        frm_buttons.pack(fill=tk.X, padx=10, pady=10)

        ttk.Button(
            frm_buttons,
            text=_("Save Excel"),
            command=lambda: self._save_excel(preview),
        ).pack(side=tk.LEFT, padx=5)

        ttk.Button(
            frm_buttons,
            text=_("Close"),
            command=preview.destroy,
        ).pack(side=tk.RIGHT, padx=5)

        # Center preview on parent
        preview.update_idletasks()
        x = self.winfo_x() + (self.winfo_width() - preview.winfo_width()) // 2
        y = self.winfo_y() + (self.winfo_height() - preview.winfo_height()) // 2
        preview.geometry(f"+{x}+{y}")

    def _save_excel(self, preview_window=None):
        """Save data to Excel with colored rows."""
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        except ImportError:
            messagebox.showerror(
                self.engine.app_title,
                _("openpyxl not installed."),
                parent=self,
            )
            return

        # Ask for save location
        filepath = filedialog.asksaveasfilename(
            parent=self,
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx")],
            initialfile=f"performance_{datetime.now().strftime('%Y%m%d')}.xlsx",
        )

        if not filepath:
            return

        self.engine.busy(self)
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Performance"

            # Define fills for coloring
            fill_high = PatternFill(start_color="FFCCCC", end_color="FFCCCC", fill_type="solid")
            fill_medium = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
            fill_good = PatternFill(start_color="CCFFCC", end_color="CCFFCC", fill_type="solid")
            fill_header = PatternFill(start_color="D0D0D0", end_color="D0D0D0", fill_type="solid")

            header_font = Font(bold=True)
            thin_border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin'),
            )

            # Headers
            headers = [
                _("Test"),
                _("Workstation"),
                _("Total"),
                _("Viol%"),
                _("Warn%"),
                _("Note%"),
                _("CV%"),
                _("Bias%"),
            ]

            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=header)
                cell.font = header_font
                cell.fill = fill_header
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="center")

            # Data rows with colors
            row_num = 2
            for item_id in self.tree.get_children():
                data = self.dict_items.get(item_id)
                values = self.tree.item(item_id, "values")

                # Determine fill based on violation level
                if data:
                    viol_pct = data.get("viol_pct", 0)
                    warn_pct = data.get("warn_pct", 0)
                    if viol_pct >= 5.0:
                        row_fill = fill_high
                    elif viol_pct >= 2.0 or warn_pct >= 10.0:
                        row_fill = fill_medium
                    else:
                        row_fill = fill_good
                else:
                    row_fill = None

                for col, val in enumerate(values, 1):
                    cell = ws.cell(row=row_num, column=col, value=val)
                    cell.border = thin_border
                    if row_fill:
                        cell.fill = row_fill
                    # Align numeric columns to right
                    if col >= 3:
                        cell.alignment = Alignment(horizontal="right")

                row_num += 1

            # Auto-fit columns
            for col in ws.columns:
                max_len = 0
                col_letter = col[0].column_letter
                for cell in col:
                    try:
                        if cell.value:
                            max_len = max(max_len, len(str(cell.value)))
                    except Exception:
                        pass
                ws.column_dimensions[col_letter].width = max_len + 2

            wb.save(filepath)

            if preview_window:
                preview_window.destroy()

            messagebox.showinfo(
                self.engine.app_title,
                f"{_('File saved')}: {filepath}",
                parent=self,
            )

        except Exception as e:
            self.engine.on_log(
                "_save_excel",
                e,
                type(e),
                __name__,
            )
            messagebox.showerror(
                self.engine.app_title,
                f"{_('Export failed')}: {e}",
                parent=self,
            )
        finally:
            self.engine.not_busy(self)

    def on_cancel(self, evt=None):
        """Close window."""
        super().on_cancel()
